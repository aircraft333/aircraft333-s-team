import pulp
import openpyxl

from config import *


def main(
    data_file=ATT1,
    template_file=TPL1,
    output_file=OUT1,
):
    data_file = resolve(data_file)
    template_file = resolve(template_file)
    output_file = resolve(output_file)

    #读取附件1数据
    df_input = pd.read_excel(data_file)
    
    price = df_input['电价'].values.astype(float)              
    load  = df_input['小区负载'].values.astype(float)            
    pv    = df_input['光伏发电预测功率'].values.astype(float)        
    
    T = N_SLOT           # 144 个 10 分钟时段
    dt = DT_H            # 10 分钟 = 1/6 小时
    
    # 储能系统参数
    E_max = SOC_MAX      # kWh，上限
    E_min = SOC_MIN      # kWh，下限
    E_0   = SOC0         # kWh，0:00 初始电量
    P_max = P_RATE       # kW，最大充放电功率
    eta   = ETA          # 充放电效率 90%
    
    #建立混合整数线性规划 (MILP)
    model = pulp.LpProblem("Microgrid_Optimal_Dispatch_Q1", pulp.LpMinimize)

    #决策变量
    P_buy = [pulp.LpVariable(f"P_buy_{t}", lowBound=0) for t in range(T)]
    P_ch  = [pulp.LpVariable(f"P_ch_{t}", lowBound=0, upBound=P_max) for t in range(T)]
    P_dis = [pulp.LpVariable(f"P_dis_{t}", lowBound=0, upBound=P_max) for t in range(T)]
    P_curt= [pulp.LpVariable(f"P_curt_{t}", lowBound=0) for t in range(T)]
    u     = [pulp.LpVariable(f"u_{t}", cat=pulp.LpBinary) for t in range(T)]
    E     = [pulp.LpVariable(f"E_{t}", lowBound=E_min, upBound=E_max) for t in range(T)]
    
    #目标函数：全天购电总费用最小
    model += pulp.lpSum([price[t] * (P_buy[t] * dt) for t in range(T)])
    
    #约束条件
    for t in range(T):
        #供需平衡
        model += P_buy[t] + pv[t] + P_dis[t] - P_ch[t] - P_curt[t] == load[t]
        
        #充放电互斥
        model += P_ch[t] <= u[t] * P_max
        model += P_dis[t] <= (1 - u[t]) * P_max
        
        #储电量状态转移
        if t == 0:
            model += E[t] == E_0 + (eta * P_ch[t] - (1.0 / eta) * P_dis[t]) * dt
        else:
            model += E[t] == E[t-1] + (eta * P_ch[t] - (1.0 / eta) * P_dis[t]) * dt
            
    #0:00 与 24:00 (t=143) 储电量相同
    model += E[T-1] == E_0
    
    # 模型求解
    solver = pulp.PULP_CBC_CMD(msg=False)
    status = model.solve(solver)
    
    if pulp.LpStatus[status] != 'Optimal':
        print("求解未达到最优")
        return
    
    #结果提取与单位换算 (kW -> kWh)
    p_buy_res = np.array([pulp.value(P_buy[t]) for t in range(T)])
    p_ch_res  = np.array([pulp.value(P_ch[t]) for t in range(T)])
    p_dis_res = np.array([pulp.value(P_dis[t]) for t in range(T)])
    e_res     = np.array([pulp.value(E[t]) for t in range(T)])
    
    q_buy = p_buy_res * dt   # 购电量 (kWh)
    q_ch  = p_ch_res * dt    # 充电量 (kWh)
    q_dis = p_dis_res * dt   # 放电量 (kWh)
    
    total_cost = np.sum(price * q_buy)
    total_buy  = np.sum(q_buy)
    
    print(f"全天总购电量: {total_buy:.4f} kWh")
    print(f"全天总购电费: {total_cost:.4f} 元")
    
    #统计 4 小时时段充放电量
    # 0:00-4:00 (0~24), 4:00-8:00 (24~48), 8:00-12:00 (48~72),
    # 12:00-16:00 (72~96), 16:00-20:00 (96~120), 20:00-24:00 (120~144)
    storage_periods = [
        ("0:00-4:00",   0,  24),
        ("4:00-8:00",  24,  48),
        ("8:00-12:00", 48,  72),
        ("12:00-16:00", 72,  96),
        ("16:00-20:00", 96, 120),
        ("20:00-24:00", 120, 144),
    ]
    
    period_ch_sums = []
    period_dis_sums = []
    for name, s, e in storage_periods:
        c_sum = np.sum(q_ch[s:e])
        d_sum = np.sum(q_dis[s:e])
        period_ch_sums.append(c_sum)
        period_dis_sums.append(d_sum)
        print(f"时段 {name:12s} | 充电量: {c_sum:10.4f} kWh | 放电量: {d_sum:10.4f} kWh")
        
    print(f"0:00 储电量: {E_0:.4f} kWh | 24:00 储电量: {e_res[-1]:.4f} kWh")
    
    #写入result1.xlsx文件
    wb = openpyxl.load_workbook(template_file)
    
    # --- 填入工作表 1: 计划购电量 ---
    ws_plan = wb["计划购电量"]
    # 模板中第 1 行为表头 ("时间段", "购电量")，数据从第 2 行到 145 行
    for idx in range(T):
        row_num = idx + 2
        # 将购电量填入第 2 列 (B列)，保留4位小数
        ws_plan.cell(row=row_num, column=2, value=round(float(q_buy[idx]), 4))
        
    # --- 填入工作表 2: 充放电量 ---
    ws_batt = wb["充放电量"]
    # 模板结构:
    # 行 1: 表头 ["时间段", "充电量", "放电量", "时刻", "储电量"]
    # 行 2-7: 6 个 4 小时时段的充放电数据
    for i in range(len(storage_periods)):
        row_num = i + 2
        ws_batt.cell(row=row_num, column=2, value=round(float(period_ch_sums[i]), 4))   # B列: 充电量
        ws_batt.cell(row=row_num, column=3, value=round(float(period_dis_sums[i]), 4))  # C列: 放电量
        
    # 填入 E列 (第5列) 的 0:00 和 24:00 储电量
    ws_batt.cell(row=2, column=5, value=round(float(E_0), 4))        # 0:00 储电量
    ws_batt.cell(row=3, column=5, value=round(float(e_res[-1]), 4))  # 24:00 储电量
    
    # 保存结果
    wb.save(output_file)

if __name__ == "__main__":
    main()