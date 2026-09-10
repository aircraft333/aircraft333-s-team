import pandas as pd
import numpy as np
import pulp
import openpyxl

def solve_microgrid_and_fill_result1(
    data_file="附件1.xlsx", 
    template_file="result1.xlsx", 
    output_file="result1.xlsx"
):
    print("=" * 60)
    print("【步骤 1】正在读取附件1数据并建立线性规划模型...")
    # 1. 读取附件1数据
    df_input = pd.read_excel(data_file)
    
    price = df_input['电价'].values.astype(float)                # 元/kWh
    load  = df_input['小区负载'].values.astype(float)              # kW
    pv    = df_input['光伏发电预测功率'].values.astype(float)          # kW
    
    T = len(price)       # 144 个 10 分钟时段
    dt = 10.0 / 60.0     # 10 分钟 = 1/6 小时
    
    # 2. 储能系统官方参数（严格依照附录1）
    E_max = 10800.0      # kWh，上限
    E_min = 1200.0       # kWh，下限
    E_0   = 6000.0       # kWh，0:00 初始电量
    P_max = 5000.0       # kW，最大充放电功率
    eta   = 0.9          # 充放电效率 90%
    
    # 3. 建立混合整数线性规划 (MILP)
    model = pulp.LpProblem("Microgrid_Optimal_Dispatch_Q1", pulp.LpMinimize)
    
    # 4. 决策变量
    P_buy = [pulp.LpVariable(f"P_buy_{t}", lowBound=0) for t in range(T)]
    P_ch  = [pulp.LpVariable(f"P_ch_{t}", lowBound=0, upBound=P_max) for t in range(T)]
    P_dis = [pulp.LpVariable(f"P_dis_{t}", lowBound=0, upBound=P_max) for t in range(T)]
    P_curt= [pulp.LpVariable(f"P_curt_{t}", lowBound=0) for t in range(T)]  # 弃光
    u     = [pulp.LpVariable(f"u_{t}", cat=pulp.LpBinary) for t in range(T)]  # 互斥变量
    E     = [pulp.LpVariable(f"E_{t}", lowBound=E_min, upBound=E_max) for t in range(T)]
    
    # 5. 目标函数：全天购电总费用最小
    model += pulp.lpSum([price[t] * (P_buy[t] * dt) for t in range(T)])
    
    # 6. 约束条件
    for t in range(T):
        # (1) 供需平衡
        model += P_buy[t] + pv[t] + P_dis[t] - P_ch[t] - P_curt[t] == load[t]
        
        # (2) 充放电互斥
        model += P_ch[t] <= u[t] * P_max
        model += P_dis[t] <= (1 - u[t]) * P_max
        
        # (3) 储电量状态转移
        if t == 0:
            model += E[t] == E_0 + (eta * P_ch[t] - (1.0 / eta) * P_dis[t]) * dt
        else:
            model += E[t] == E[t-1] + (eta * P_ch[t] - (1.0 / eta) * P_dis[t]) * dt
            
    # (4) 0:00 与 24:00 (t=143) 储电量相同
    model += E[T-1] == E_0
    
    # 7. 模型求解
    print("【步骤 2】正在调用 CBC 求解器求解最优调度方案...")
    solver = pulp.PULP_CBC_CMD(msg=False)
    status = model.solve(solver)
    
    if pulp.LpStatus[status] != 'Optimal':
        print("❌ 求解未达到最优，请检查数据与约束！")
        return
    
    # 8. 结果提取与单位换算 (kW -> kWh)
    p_buy_res = np.array([pulp.value(P_buy[t]) for t in range(T)])
    p_ch_res  = np.array([pulp.value(P_ch[t]) for t in range(T)])
    p_dis_res = np.array([pulp.value(P_dis[t]) for t in range(T)])
    e_res     = np.array([pulp.value(E[t]) for t in range(T)])
    
    q_buy = p_buy_res * dt   # 购电量 (kWh)
    q_ch  = p_ch_res * dt    # 充电量 (kWh)
    q_dis = p_dis_res * dt   # 放电量 (kWh)
    
    total_cost = np.sum(price * q_buy)
    total_buy  = np.sum(q_buy)
    
    print("\n" + "=" * 60)
    print(f"✅ 全局最优解求解成功！")
    print(f"全天总购电量: {total_buy:.4f} kWh")
    print(f"全天总购电费: {total_cost:.4f} 元")
    print("=" * 60)
    
    # 9. 统计 4 小时时段充放电量
    # 0:00-4:00 (0~24), 4:00-8:00 (24~48), 8:00-12:00 (48~72),
    # 12:00-16:00 (72~96), 16:00-20:00 (96~120), 20:00-24:00 (120~144)
    storage_periods = [
        ("0:00-4:00",   0,  24),
        ("4:00-8:00",  24,  48),
        ("8:00-12:00", 48,  72),
        ("12:00-16:00", 72,  96),
        ("16:00-20:00", 96, 120),
        ("20:00-24:00", 120, 144),
    ]
    
    period_ch_sums = []
    period_dis_sums = []
    for name, s, e in storage_periods:
        c_sum = np.sum(q_ch[s:e])
        d_sum = np.sum(q_dis[s:e])
        period_ch_sums.append(c_sum)
        period_dis_sums.append(d_sum)
        print(f"时段 {name:12s} | 充电量: {c_sum:10.4f} kWh | 放电量: {d_sum:10.4f} kWh")
        
    print(f"0:00 储电量: {E_0:.4f} kWh | 24:00 储电量: {e_res[-1]:.4f} kWh")
    
    # 10. 【核心】直接写入官方 result1.xlsx 模板文件
    print(f"\n【步骤 3】正在将计算结果填入 {output_file} ...")
    wb = openpyxl.load_workbook(template_file)
    
    # --- 填入工作表 1: 计划购电量 ---
    ws_plan = wb["计划购电量"]
    # 模板中第 1 行为表头 ("时间段", "购电量")，数据从第 2 行到 145 行
    for idx in range(T):
        row_num = idx + 2
        # 将购电量填入第 2 列 (B列)，保留4位小数
        ws_plan.cell(row=row_num, column=2, value=round(float(q_buy[idx]), 4))
        
    # --- 填入工作表 2: 充放电量 ---
    ws_batt = wb["充放电量"]
    # 模板结构:
    # 行 1: 表头 ["时间段", "充电量", "放电量", "时刻", "储电量"]
    # 行 2-7: 6 个 4 小时时段的充放电数据
    for i in range(len(storage_periods)):
        row_num = i + 2
        ws_batt.cell(row=row_num, column=2, value=round(float(period_ch_sums[i]), 4))   # B列: 充电量
        ws_batt.cell(row=row_num, column=3, value=round(float(period_dis_sums[i]), 4))  # C列: 放电量
        
    # 填入 E列 (第5列) 的 0:00 和 24:00 储电量
    ws_batt.cell(row=2, column=5, value=round(float(E_0), 4))        # 0:00 储电量
    ws_batt.cell(row=3, column=5, value=round(float(e_res[-1]), 4))  # 24:00 储电量
    
    # 保存结果
    wb.save(output_file)
    print(f"🎉 全部结果已成功填入并保存至文件: {output_file}")

if __name__ == "__main__":
    # 使用前请确保已安装所需库:
    # pip install pulp openpyxl pandas numpy
    solve_microgrid_and_fill_result1(
        data_file="附件1.xlsx",
        template_file="result1.xlsx",
        output_file="result1.xlsx"
    )