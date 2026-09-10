# -*- coding: utf-8 -*-
"""
Q1 微电网优化调度 —— 含对照实验（证明模型有效性）
无 tier 依赖：根据 price 自动划分峰/平/谷
"""

import pulp
import openpyxl
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

from config import *   # ATT1, TPL1, OUT1, N_SLOT, DT_H,
                       # SOC_MAX, SOC_MIN, SOC0, P_RATE, ETA

plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei"]
plt.rcParams["axes.unicode_minus"] = False

# ---- 峰谷划分阈值（可调） ----
PEAK_Q = 0.70    # 价格 >= 70% 分位数 → 峰
VALLEY_Q = 0.30  # 价格 <= 30% 分位数 → 谷
# 其余为平


def build_tier(price, peak_q=PEAK_Q, valley_q=VALLEY_Q):
    """根据价格分位数划分：0=谷, 1=平, 2=峰"""
    hi = np.quantile(price, peak_q)
    lo = np.quantile(price, valley_q)
    tier = np.ones(len(price), dtype=int)      # 默认平
    tier[price <= lo] = 0                      # 谷
    tier[price >= hi] = 2                      # 峰
    return tier


# =========================================================
# 基准策略
# =========================================================
def baseline_no_storage(price, load, pv, dt, T):
    """基准1：不装储能。光伏先自用，不足买电，富余弃掉。"""
    net = load - pv
    q_buy = np.maximum(net, 0.0) * dt
    q_curt = np.maximum(-net, 0.0) * dt
    cost = float(np.sum(price * q_buy))
    return dict(cost=cost, q_buy=q_buy, q_curt=q_curt)


def baseline_rule_based(price, load, pv, dt, T, tier,
                        E_max, E_min, E_0, P_max, eta):
    """基准2：规则型——谷段满充，峰段满放，平段不动。"""
    E = E_0
    q_buy = np.zeros(T)
    q_curt = np.zeros(T)
    q_ch = np.zeros(T)
    q_dis = np.zeros(T)
    e_traj = np.zeros(T)
    for t in range(T):
        net = load[t] - pv[t]
        if tier[t] == 0:                        # 谷：尽量充
            ch = min(P_max, (E_max - E) / (eta * dt))
            ch = max(ch, 0.0)
            E += eta * ch * dt
            net += ch
            q_ch[t] = ch * dt
        elif tier[t] == 2:                      # 峰：尽量放
            dis = min(P_max, (E - E_min) * eta / dt)
            dis = max(dis, 0.0)
            E -= dis / eta * dt
            net -= dis
            q_dis[t] = dis * dt
        # 平段不动
        if net > 0:
            q_buy[t] = net * dt
        else:
            q_curt[t] = -net * dt
        e_traj[t] = E
    cost = float(np.sum(price * q_buy))
    return dict(cost=cost, q_buy=q_buy, q_curt=q_curt,
                q_ch=q_ch, q_dis=q_dis, e_traj=e_traj)


# =========================================================
# 主流程
# =========================================================
def main(
    data_file=ATT1,
    template_file=TPL1,
    output_file=OUT1,
):
    data_file = resolve(data_file)
    template_file = resolve(template_file)
    output_file = resolve(output_file)

    # ---------- 读取数据 ----------
    df_input = pd.read_excel(data_file)
    price = df_input['电价'].values.astype(float)
    load  = df_input['小区负载'].values.astype(float)
    pv    = df_input['光伏发电预测功率'].values.astype(float)

    T  = N_SLOT
    dt = DT_H

    E_max = SOC_MAX
    E_min = SOC_MIN
    E_0   = SOC0
    P_max = P_RATE
    eta   = ETA

    # ---------- 自动划分峰/平/谷 ----------
    tier = build_tier(price)
    n_v = int(np.sum(tier == 0))
    n_f = int(np.sum(tier == 1))
    n_p = int(np.sum(tier == 2))
    print(f"峰谷划分：谷 {n_v} 段 | 平 {n_f} 段 | 峰 {n_p} 段")
    print(f"价格分位：谷阈值 {np.quantile(price, VALLEY_Q):.4f} | "
          f"峰阈值 {np.quantile(price, PEAK_Q):.4f}")

    # ---------- 建立 MILP ----------
    model = pulp.LpProblem("Microgrid_Optimal_Dispatch_Q1", pulp.LpMinimize)

    P_buy = [pulp.LpVariable(f"P_buy_{t}", lowBound=0) for t in range(T)]
    P_ch  = [pulp.LpVariable(f"P_ch_{t}",  lowBound=0, upBound=P_max) for t in range(T)]
    P_dis = [pulp.LpVariable(f"P_dis_{t}", lowBound=0, upBound=P_max) for t in range(T)]
    P_curt= [pulp.LpVariable(f"P_curt_{t}",lowBound=0) for t in range(T)]
    u     = [pulp.LpVariable(f"u_{t}", cat=pulp.LpBinary) for t in range(T)]
    E     = [pulp.LpVariable(f"E_{t}", lowBound=E_min, upBound=E_max) for t in range(T)]

    model += pulp.lpSum([price[t] * (P_buy[t] * dt) for t in range(T)])

    for t in range(T):
        model += P_buy[t] + pv[t] + P_dis[t] - P_ch[t] - P_curt[t] == load[t]
        model += P_ch[t]  <= u[t] * P_max
        model += P_dis[t] <= (1 - u[t]) * P_max
        model += P_curt[t] <= pv[t]
        if t == 0:
            model += E[t] == E_0 + (eta * P_ch[t] - (1.0 / eta) * P_dis[t]) * dt
        else:
            model += E[t] == E[t-1] + (eta * P_ch[t] - (1.0 / eta) * P_dis[t]) * dt

    model += E[T-1] == E_0

    solver = pulp.PULP_CBC_CMD(msg=False)
    status = model.solve(solver)
    if pulp.LpStatus[status] != 'Optimal':
        print("求解未达到最优"); return

    # ---------- 提取优化结果 ----------
    p_buy_res = np.array([pulp.value(P_buy[t]) for t in range(T)])
    p_ch_res  = np.array([pulp.value(P_ch[t])  for t in range(T)])
    p_dis_res = np.array([pulp.value(P_dis[t]) for t in range(T)])
    p_curt_res= np.array([pulp.value(P_curt[t])for t in range(T)])
    e_res     = np.array([pulp.value(E[t])     for t in range(T)])

    q_buy  = p_buy_res  * dt
    q_ch   = p_ch_res   * dt
    q_dis  = p_dis_res  * dt
    q_curt = p_curt_res * dt

    total_cost = float(np.sum(price * q_buy))
    total_buy  = float(np.sum(q_buy))

    print(f"全天总购电量: {total_buy:.4f} kWh")
    print(f"全天总购电费: {total_cost:.4f} 元")

    # ---------- 6 个 4 小时时段统计 ----------
    storage_periods = [
        ("0:00-4:00",   0,  24),
        ("4:00-8:00",  24,  48),
        ("8:00-12:00", 48,  72),
        ("12:00-16:00", 72,  96),
        ("16:00-20:00", 96, 120),
        ("20:00-24:00", 120, 144),
    ]
    period_ch_sums, period_dis_sums = [], []
    for name, s, e in storage_periods:
        c_sum = float(np.sum(q_ch[s:e]))
        d_sum = float(np.sum(q_dis[s:e]))
        period_ch_sums.append(c_sum)
        period_dis_sums.append(d_sum)
        print(f"时段 {name:12s} | 充电量: {c_sum:10.4f} kWh | 放电量: {d_sum:10.4f} kWh")
    print(f"0:00 储电量: {E_0:.4f} kWh | 24:00 储电量: {e_res[-1]:.4f} kWh")

    # =====================================================
    # 对照实验
    # =====================================================
    b1 = baseline_no_storage(price, load, pv, dt, T)
    b2 = baseline_rule_based(price, load, pv, dt, T, tier,
                             E_max, E_min, E_0, P_max, eta)

    def pv_util_rate(q_curt_):
        pv_total = float(np.sum(pv) * dt)
        return 1.0 - float(np.sum(q_curt_)) / pv_total if pv_total > 0 else 0.0

    def peak_buy_ratio(q_buy_):
        tot = float(np.sum(q_buy_))
        if tot <= 0: return 0.0
        return float(np.sum(q_buy_[tier == 2])) / tot

    metrics = {
        "无储能": {
            "cost": b1["cost"],
            "buy": float(np.sum(b1["q_buy"])),
            "pv_util": pv_util_rate(b1["q_curt"]),
            "peak_ratio": peak_buy_ratio(b1["q_buy"]),
            "curt": float(np.sum(b1["q_curt"])),
        },
        "规则谷充峰放": {
            "cost": b2["cost"],
            "buy": float(np.sum(b2["q_buy"])),
            "pv_util": pv_util_rate(b2["q_curt"]),
            "peak_ratio": peak_buy_ratio(b2["q_buy"]),
            "curt": float(np.sum(b2["q_curt"])),
        },
        "优化模型": {
            "cost": total_cost,
            "buy": total_buy,
            "pv_util": pv_util_rate(q_curt),
            "peak_ratio": peak_buy_ratio(q_buy),
            "curt": float(np.sum(q_curt)),
        },
    }

    print("\n========== 对照实验结果 ==========")
    print(f"{'策略':<14}{'购电费(元)':>14}{'购电量(kWh)':>16}"
          f"{'光伏消纳率':>12}{'峰段占比':>10}{'弃光(kWh)':>12}")
    for k, v in metrics.items():
        print(f"{k:<14}{v['cost']:>14.2f}{v['buy']:>16.2f}"
              f"{v['pv_util']*100:>11.1f}%{v['peak_ratio']*100:>9.1f}%"
              f"{v['curt']:>12.2f}")

    save1 = b1["cost"] - total_cost
    save2 = b2["cost"] - total_cost
    print(f"\n相比【无储能】节省: {save1:.2f} 元 ({save1/b1['cost']*100:.2f}%)")
    print(f"相比【规则型】节省: {save2:.2f} 元 ({save2/b2['cost']*100:.2f}%)")

    # =====================================================
    # 可视化
    # =====================================================
    hours = np.arange(T) * dt

    # 图A：三种策略购电功率曲线
    fig, ax = plt.subplots(figsize=(12, 4.5))
    ax.plot(hours, b1["q_buy"]/dt, label="无储能", lw=1.5, alpha=0.8)
    ax.plot(hours, b2["q_buy"]/dt, label="规则谷充峰放", lw=1.5, alpha=0.8)
    ax.plot(hours, q_buy/dt,       label="优化模型", lw=2.0)
    ax.set_xlabel("时间 (h)"); ax.set_ylabel("购电功率 (kW)")
    ax.set_title("图A  三种策略购电功率对比", fontweight="bold")
    ax.legend(); ax.grid(alpha=0.3); ax.margins(x=0)
    fig.tight_layout(); save_fig(fig, "figA_购电对比.png")

    # 图B：费用柱状图
    fig, ax = plt.subplots(figsize=(6, 4.5))
    names = list(metrics.keys())
    costs = [metrics[n]["cost"] for n in names]
    bars = ax.bar(names, costs, color=["#999999", "#1f77b4", "#d62728"])
    for bar, c in zip(bars, costs):
        ax.text(bar.get_x()+bar.get_width()/2, c, f"{c:.0f}",
                ha="center", va="bottom", fontsize=10)
    ax.set_ylabel("全天购电费 (元)")
    ax.set_title("图B  三种策略总费用对比", fontweight="bold")
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout(); save_fig(fig, "figB_费用对比.png")

    # 图C：SOC 轨迹
    fig, ax = plt.subplots(figsize=(12, 4))
    ax.plot(hours, e_res, label="优化 SOC", lw=2)
    ax.plot(hours, b2["e_traj"], label="规则 SOC", lw=1.5, ls="--")
    ax.axhline(E_max, color="r", ls=":", lw=1, label="SOC 上限")
    ax.axhline(E_min, color="g", ls=":", lw=1, label="SOC 下限")
    ax.set_xlabel("时间 (h)"); ax.set_ylabel("储电量 (kWh)")
    ax.set_title("图C  储能 SOC 轨迹对比", fontweight="bold")
    ax.legend(); ax.grid(alpha=0.3); ax.margins(x=0)
    fig.tight_layout(); save_fig(fig, "figC_SOC对比.png")

    # =====================================================
    # 写入 Excel
    # =====================================================
    wb = openpyxl.load_workbook(template_file)

    ws_plan = wb["计划购电量"]
    for idx in range(T):
        ws_plan.cell(row=idx+2, column=2, value=round(float(q_buy[idx]), 4))

    ws_batt = wb["充放电量"]
    for i in range(len(storage_periods)):
        ws_batt.cell(row=i+2, column=2, value=round(float(period_ch_sums[i]), 4))
        ws_batt.cell(row=i+2, column=3, value=round(float(period_dis_sums[i]), 4))
    ws_batt.cell(row=2, column=5, value=round(float(E_0), 4))
    ws_batt.cell(row=3, column=5, value=round(float(e_res[-1]), 4))

    if "对照实验" in wb.sheetnames:
        wb.remove(wb["对照实验"])
    ws_cmp = wb.create_sheet("对照实验")
    ws_cmp.append(["策略", "购电费(元)", "购电量(kWh)",
                   "光伏消纳率", "峰段购电占比", "弃光量(kWh)"])
    for k, v in metrics.items():
        ws_cmp.append([k, round(v["cost"], 2), round(v["buy"], 2),
                       f"{v['pv_util']*100:.1f}%",
                       f"{v['peak_ratio']*100:.1f}%",
                       round(v["curt"], 2)])
    ws_cmp.append([])
    ws_cmp.append(["相比无储能节省(元)", round(save1, 2)])
    ws_cmp.append(["相比规则型节省(元)", round(save2, 2)])

    wb.save(output_file)
    print(f"\n结果已写入: {output_file}")


if __name__ == "__main__":
    main()