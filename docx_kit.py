# -*- coding: utf-8 -*-
"""《模型与工作流》docx 生成工具箱

把「中文正文 + LaTeX 源码 + 公式渲染图」三件套的构建逻辑集中在这里，
`make_q3_docx.py` / `make_q4_docx.py` 只负责写各问的内容。

公式渲染管线（本机 TeX Live 2026）：
    latex → .dvi → dvipng -T tight -D 200 -bg Transparent → 透明 PNG
注意不要用 preview 宏包的 tightpage（会报 "No pages of output"），
普通单页 + dvipng 的 -T tight 已经能自动裁边。
"""
import hashlib
import os
import subprocess

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DPI = 200                       # 渲染分辨率，用于把像素换算成英寸
MAX_W = 6.2                     # 公式图片最大宽度（英寸）
GRAY = RGBColor(0x60, 0x60, 0x60)
BLUE = RGBColor(0x1F, 0x49, 0x7D)

TEX_TPL = ("\\documentclass[12pt]{article}\n"
           "\\pagestyle{empty}\n"
           "\\usepackage{amsmath,amssymb}\n"
           "\\begin{document}\n"
           "\\noindent$%s$\n"
           "\\end{document}\n")


def render_tex(formula, eqdir):
    """把 LaTeX 公式渲染成透明 PNG，返回文件路径（按内容哈希缓存）"""
    os.makedirs(eqdir, exist_ok=True)
    tag = hashlib.md5(formula.encode("utf-8")).hexdigest()[:12]
    png = os.path.join(eqdir, tag + ".png")
    if os.path.exists(png):
        return png
    with open(os.path.join(eqdir, "_tmp.tex"), "w", encoding="utf-8") as f:
        f.write(TEX_TPL % formula)
    r = subprocess.run(["latex", "-interaction=nonstopmode", "_tmp.tex"],
                       cwd=eqdir, capture_output=True, text=True)
    if not os.path.exists(os.path.join(eqdir, "_tmp.dvi")):
        raise RuntimeError("LaTeX 渲染失败：%s\n%s" % (formula[:90], r.stdout[-800:]))
    r2 = subprocess.run(["dvipng", "-T", "tight", "-D", str(DPI),
                         "-bg", "Transparent", "-o", png, "_tmp.dvi"],
                        cwd=eqdir, capture_output=True, text=True)
    if r2.returncode != 0 or not os.path.exists(png):
        raise RuntimeError("dvipng 转换失败：%s" % r2.stderr[-400:])
    return png


class Doc:
    """按《问题X 模型与工作流》的统一排版逐步构建文档"""

    def __init__(self, title_text, eqdir):
        self.eqdir = eqdir
        self.doc = Document()
        n = self.doc.styles["Normal"]
        n.font.name = "Times New Roman"
        n.font.size = Pt(10.5)
        n._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        n.paragraph_format.space_after = Pt(3)
        n.paragraph_format.line_spacing = 1.25
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        self._run(p, title_text, "黑体", 17, True)

    # ---------- 内部 ----------
    def _p(self):
        return self.doc.add_paragraph()

    @staticmethod
    def _run(p, text, name="宋体", size=10.5, bold=False, color=None, mono=False):
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = "Consolas" if mono else name
        r._element.rPr.rFonts.set(qn("w:eastAsia"), name)
        if color:
            r.font.color.rgb = color
        return r

    def _rich(self, p, text, size=10.5, bold=False, color=None):
        """按 **…** 切分成粗体段，支持正文里的局部强调"""
        for i, seg in enumerate(text.split("**")):
            if seg:
                self._run(p, seg, size=size, bold=(bold or i % 2 == 1), color=color)

    # ---------- 对外 ----------
    def h1(self, text):
        p = self._p()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        self._run(p, text, "黑体", 13.5, True)

    def h2(self, text):
        p = self._p()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(5)
        self._run(p, text, "黑体", 11.5, True)

    def para(self, text, indent=True, bold=False):
        p = self._p()
        if indent:
            p.paragraph_format.first_line_indent = Pt(21)
        self._rich(p, text, size=10.5, bold=bold)

    def item(self, text, marker="•"):
        p = self._p()
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.space_after = Pt(2)
        self._run(p, marker + " ", size=10.5, bold=True)
        self._rich(p, text, size=10.5)

    def src(self, latex):
        """LaTeX 源码（等宽小字、灰色），便于直接复制进论文"""
        p = self._p()
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.space_after = Pt(2)
        self._run(p, latex, size=8.5, mono=True, color=GRAY)

    def eq(self, latex, show_src=True):
        """插入公式：先给 LaTeX 源码，再给渲染图"""
        if show_src:
            self.src(latex)
        png = render_tex(latex, self.eqdir)
        w, _h = Image.open(png).size
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        p.add_run().add_picture(png, width=Inches(min(w / DPI, MAX_W)))

    def note(self, text):
        p = self._p()
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.space_after = Pt(4)
        self._run(p, "注：", size=9.5, bold=True, color=BLUE)
        self._run(p, text, size=9.5, color=BLUE)

    def save(self, path):
        """保存文档；被 Word 占用时给出可操作的中文报错"""
        try:
            self.doc.save(path)
        except PermissionError:
            lock = os.path.join(os.path.dirname(path),
                                "~$" + os.path.basename(path))
            raise PermissionError(
                "无法写入 {0}：文件正被 Word 等程序占用。\n"
                "  请先关闭它再重跑；锁文件 = {1}".format(
                    os.path.basename(path), lock)) from None
        return path
