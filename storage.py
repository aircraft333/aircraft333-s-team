# -*- coding: utf-8 -*-
"""
2026 高教社杯 C 题 · 问题 1 求解、表格填报与10分钟储电量导出
=====================================================================
严格依照 config.py 统一接口、参数名称、数据加载及路径规范
"""
import pulp
import openpyxl

from config import *

def main(
    data_file=ATT1,
    template_file=TPL1,
    output_file=OUT1,
    storage_file="storage.xlsx",
):
    # 路径解析（统一使用 resolve 保证跨目录安全）
    data_path = resolve(data_file)
    template_path = resolve(template_file)
    if not os.path.exists(template_path):
        template_path = resolve(output_file)  # 回退到根目录模板
    output_path = resolve(output_file)
    storage_path = resolve(storage_file)      # storage.xlsx 绝对路径

    rule("【问题 1】微网日前计划购电策略优化与填报")

    # 1. 使用 config.py 标准数据读取与数组提取接口
    df_input = load_att1(data_path)
    arrs = att1_arrays(df_input)
    
    price = arrs["price"]   # 分时电价 (元/kWh)
    load = arrs["load"]     # 小区负载 (kW)
    pv = arrs["pv"]         # 光伏预测功率 (kW)
    
    # 2. 建立混合整数线性规划 (MILP) 模型
    model = pulp.LpProblem("Microgrid_Optimal_Dispatch_Q1", pulp.LpMinimize)
    
    # 3. 决策变量声明（直接使用 config.py 命名）
    P_buy = [pulp.LpVariable(f"P_buy_{t}", lowBound=0) for t in range(N_SLOT)]
    P_ch  = [pulp.LpVariable(f"P_ch_{t}", lowBound=0, upBound=P_RATE) for t in range(N_SLOT)]
    P_dis = [pulp.LpVariable(f"P_dis_{t}", lowBound=0, upBound=P_RATE) for t in range(N_SLOT)]
    P_curt= [pulp.LpVariable(f"P_curt_{t}", lowBound=0) for t in range(N_SLOT)]
    u     = [pulp.LpVariable(f"u_{t}", cat=pulp.LpBinary) for t in range(N_SLOT)]
    E     = [pulp.LpVariable(f"E_{t}", lowBound=SOC_MIN, upBound=SOC_MAX) for t in range(N_SLOT)]
    
    # 4. 目标函数：全天购电成本最小化
    model += pulp.lpSum([price[t] * (P_buy[t] * DT_H) for t in range(N_SLOT)])
    
    # 5. 约束条件
    for t in range(N_SLOT):
        # (1) 微网供需功率平衡
        model += P_buy[t] + pv[t] + P_dis[t] - P_ch[t] - P_curt[t] == load[t]
        
        # (2) 充放电功率互斥
        model += P_ch[t] <= u[t] * P_RATE
        model += P_dis[t] <= (1 - u[t]) * P_RATE
        
        # (3) 储能量动态转移方程
        if t == 0:
            model += E[t] == SOC0 + (ETA * P_ch[t] - (1.0 / ETA) * P_dis[t]) * DT_H
        else:
            model += E[t] == E[t-1] + (ETA * P_ch[t] - (1.0 / ETA) * P_dis[t]) * DT_H
            
    # (4) 0:00 与 24:00 (t=143 结束) 储电量相同
    model += E[N_SLOT - 1] == SOC0
    
    # 6. 模型求解
    log("正在调用 CBC 求解器求解...")
    solver = pulp.PULP_CBC_CMD(msg=False)
    status = model.solve(solver)
    
    if pulp.LpStatus[status] != 'Optimal':
        log("❌ 求解未达到最优，请检查数据与约束！")
        return
    
    # 7. 提取变量解并转换为电量 (kWh)
    p_buy_res = np.array([pulp.value(P_buy[t]) for t in range(N_SLOT)])
    p_ch_res  = np.array([pulp.value(P_ch[t]) for t in range(N_SLOT)])
    p_dis_res = np.array([pulp.value(P_dis[t]) for t in range(N_SLOT)])
    e_res     = np.array([pulp.value(E[t]) for t in range(N_SLOT)])
    
    q_buy = p_buy_res * DT_H
    q_ch  = p_ch_res * DT_H
    q_dis = p_dis_res * DT_H
    
    total_cost = np.sum(price * q_buy)
    total_buy  = np.sum(q_buy)
    
    rule("求解结果汇总")
    log(f"全天总购电量: {total_buy:.4f} kWh")
    log(f"全天总购电费: {total_cost:.4f} 元")
    
    # 8. 统计 4 小时时段充放电量（每 4 小时 = 24 个时段）
    storage_periods = [
        ("0:00-4:00",   0,  24),
        ("4:00-8:00",  24,  48),
        ("8:00-12:00", 48,  72),
        ("12:00-16:00", 72,  96),
        ("16:00-20:00", 96, 120),
        ("20:00-24:00", 120, 144),
    ]
    
    period_ch_sums = []
    period_dis_sums = []
    for name, s, e in storage_periods:
        c_sum = float(np.sum(q_ch[s:e]))
        d_sum = float(np.sum(q_dis[s:e]))
        period_ch_sums.append(c_sum)
        period_dis_sums.append(d_sum)
        log(f"时段 {name:12s} | 充电量: {c_sum:10.4f} kWh | 放电量: {d_sum:10.4f} kWh")
        
    log(f"0:00 储电量: {SOC0:.4f} kWh | 24:00 储电量: {e_res[-1]:.4f} kWh")
    
    # 9. 填入 result1.xlsx 模板
    log(f"\n正在将数据写入模板并保存至 -> {output_path}")
    wb = openpyxl.load_workbook(template_path)
    
    # --- 工作表 1: 计划购电量 ---
    ws_plan = wb["计划购电量"]
    for idx in range(N_SLOT):
        row_num = idx + 2
        ws_plan.cell(row=row_num, column=2, value=round(float(q_buy[idx]), 4))
        
    # --- 工作表 2: 充放电量 ---
    ws_batt = wb["充放电量"]
    for i in range(len(storage_periods)):
        row_num = i + 2
        ws_batt.cell(row=row_num, column=2, value=round(period_ch_sums[i], 4))
        ws_batt.cell(row=row_num, column=3, value=round(period_dis_sums[i], 4))
        
    ws_batt.cell(row=2, column=5, value=round(float(SOC0), 4))
    ws_batt.cell(row=3, column=5, value=round(float(e_res[-1]), 4))
    
    wb.save(output_path)
    log(f"🎉 result1.xlsx 表格填报已完成并成功保存！")

    # 10. 【新增】将每 10 分钟时段及其对应的储电量写入 storage.xlsx
    log(f"正在导出实时（10分钟间隔）储电量至 -> {storage_path}")
    
    # 构建时段字符串，例如 "00:00-00:10"
    time_intervals = [
        f"{fmt(df_input['分钟'].iloc[i] - 10)}-{fmt(df_input['分钟'].iloc[i])}"
        for i in range(N_SLOT)
    ]
    
    df_storage = pd.DataFrame({
        "时段序号": range(1, N_SLOT + 1),
        "时间段": time_intervals,
        "时刻(右端点)": df_input["时间"],
        "储电量(kWh)": np.round(e_res, 4),
        "充电功率(kW)": np.round(p_ch_res, 4),
        "放电功率(kW)": np.round(p_dis_res, 4)
    })
    
    # 写入 storage.xlsx
    df_storage.to_excel(storage_path, index=False)
    log(f"🎉 storage.xlsx 导出成功（共 {len(df_storage)} 行数据）！")

if __name__ == "__main__":
    main()